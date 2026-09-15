"""
Generate the mock CRM workbook that Layer 3 (the UiPath robot) acts on.

Standing in for a production order database, per the backend-simulation plan:
an Excel workbook with the same fields a real CRM would expose, reachable
through the REST-style wrapper in rpa/crm.py.

The data is deliberately constructed, not random. Every demo scenario needs a
row that exercises it, so the set contains orders that auto-resolve, orders
that must escalate on the sensitivity cap, orders that are refund-ineligible,
and orders too far along to redirect. Regenerate any time with:

    python scripts/build_mock_crm.py

That resets AuditLog and EscalationQueue to empty, which is exactly what you
want between rehearsal runs and the real demo.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import config  # noqa: E402

ORDER_COLUMNS = [
    "OrderID", "CustomerName", "CustomerEmail", "Phone", "Category",
    "Product", "OrderDate", "Amount", "Status", "RefundEligible",
    "DeliveryAddress", "TrackingNo",
]

AUDIT_COLUMNS = [
    "Timestamp", "OrderID", "Intent", "Confidence", "Engine", "Action",
    "Result", "ReferenceNo", "HandledBy",
]

ESCALATION_COLUMNS = [
    "Timestamp", "OrderID", "Intent", "Confidence", "Entities", "Reason",
    "SuggestedResolution", "CustomerMessage", "Status", "AssignedTo",
]

# OrderID, Customer, Email, Phone, Category, Product, Date, Amount, Status,
# RefundEligible, Address, Tracking
ORDERS = [
    # -- Television ---------------------------------------------------------
    (4501, "Rahul Sharma", "rahul.sharma@example.com", "9820011223",
     "Television", "Sony Bravia 43\" 4K LED", "2026-08-28", 48990.00,
     "Delivered", "N", "12 Hill View, Andheri West, Mumbai 400058", "TRK45010"),
    (4502, "Priya Nair", "priya.nair@example.com", "9845112233",
     "Television", "Samsung Crystal 50\" UHD", "2026-09-05", 54990.00,
     "Shipped", "Y", "7B Green Acres, Koramangala, Bengaluru 560034", "TRK45020"),
    (4503, "Amit Deshpande", "amit.d@example.com", "9730223344",
     "Television", "LG 32\" HD Smart TV", "2026-09-11", 16499.00,
     "Processing", "Y", "402 Sunshine CHS, Kothrud, Pune 411038", ""),
    (4504, "Sneha Iyer", "sneha.iyer@example.com", "9884334455",
     "Television", "Mi 43\" Full HD Android TV", "2026-09-09", 24999.00,
     "Out for Delivery", "Y", "23 Lake Road, Adyar, Chennai 600020", "TRK45040"),

    # -- Washing Machine ----------------------------------------------------
    (4505, "Vikram Singh", "vikram.singh@example.com", "9811445566",
     "Washing Machine", "IFB 7kg Front Load", "2026-09-02", 32490.00,
     "Delivered", "Y", "D-14 Rajouri Garden, New Delhi 110027", "TRK45050"),
    (4506, "Anjali Mehta", "anjali.mehta@example.com", "9825556677",
     "Washing Machine", "LG 6.5kg Top Load", "2026-09-12", 18990.00,
     "Processing", "Y", "55 Satellite Road, Ahmedabad 380015", ""),
    (4507, "Karthik Reddy", "karthik.r@example.com", "9866667788",
     "Washing Machine", "Bosch 8kg Front Load", "2026-08-21", 41990.00,
     "Delivered", "N", "18 Jubilee Hills, Hyderabad 500033", "TRK45070"),
    (4508, "Meera Joshi", "meera.joshi@example.com", "9822778899",
     "Washing Machine", "Whirlpool 7kg Semi Automatic", "2026-09-10", 12499.00,
     "Shipped", "Y", "9 Camp Area, Nagpur 440001", "TRK45080"),

    # -- Mixer Grinder ------------------------------------------------------
    (4509, "Rohit Verma", "rohit.verma@example.com", "9812889900",
     "Mixer Grinder", "Philips HL7756 750W", "2026-09-08", 4299.00,
     "Delivered", "Y", "301 Sector 21, Noida 201301", "TRK45090"),
    (4510, "Divya Menon", "divya.menon@example.com", "9847990011",
     "Mixer Grinder", "Preethi Zodiac 750W", "2026-09-13", 7499.00,
     "Processing", "Y", "14 MG Road, Kochi 682016", ""),
    (4511, "Sanjay Gupta", "sanjay.gupta@example.com", "9831001122",
     "Mixer Grinder", "Bajaj Rex 500W", "2026-08-30", 2199.00,
     "Delivered", "N", "27 Park Street, Kolkata 700016", "TRK45110"),
    (4512, "Nisha Rao", "nisha.rao@example.com", "9880112233",
     "Mixer Grinder", "Butterfly Smart 750W", "2026-09-12", 3499.00,
     "Cancelled", "Y", "6 Vijayanagar, Mysuru 570017", ""),

    # -- Refrigerator -------------------------------------------------------
    (4513, "Arjun Pillai", "arjun.pillai@example.com", "9895223344",
     "Refrigerator", "Samsung 253L Double Door", "2026-09-04", 27990.00,
     "Shipped", "Y", "88 Marine Drive, Thiruvananthapuram 695001", "TRK45130"),
    (4514, "Pooja Bhatt", "pooja.bhatt@example.com", "9820334455",
     "Refrigerator", "LG 190L Single Door", "2026-09-14", 15499.00,
     "Processing", "Y", "11 Linking Road, Bandra, Mumbai 400050", ""),
    (4515, "Manish Agarwal", "manish.a@example.com", "9829445566",
     "Refrigerator", "Whirlpool 340L Frost Free", "2026-08-25", 38990.00,
     "Delivered", "N", "3 Civil Lines, Jaipur 302006", "TRK45150"),
    (4516, "Lakshmi Subramanian", "lakshmi.s@example.com", "9840556677",
     "Refrigerator", "Haier 596L Side by Side", "2026-09-06", 72990.00,
     "Delivered", "Y", "45 Anna Nagar, Chennai 600040", "TRK45160"),

    # -- Air Conditioner ----------------------------------------------------
    (4517, "Farhan Qureshi", "farhan.q@example.com", "9819667788",
     "Air Conditioner", "Voltas 1.5T 3 Star Split", "2026-09-07", 34990.00,
     "Delivered", "Y", "62 Hazratganj, Lucknow 226001", "TRK45170"),
    (4518, "Kavita Desai", "kavita.desai@example.com", "9825778899",
     "Air Conditioner", "Blue Star 1T 5 Star Split", "2026-09-13", 38490.00,
     "Processing", "Y", "29 Alkapuri, Vadodara 390007", ""),
    (4519, "Suresh Kumar", "suresh.kumar@example.com", "9844889900",
     "Air Conditioner", "Daikin 1.5T Inverter Split", "2026-09-01", 45990.00,
     "Out for Delivery", "Y", "7 Race Course, Coimbatore 641018", "TRK45190"),
    (4520, "Ritu Chawla", "ritu.chawla@example.com", "9810990011",
     "Air Conditioner", "Lloyd 1.5T Window AC", "2026-08-27", 26490.00,
     "Refunded", "N", "16 Model Town, Ludhiana 141002", "TRK45200"),

    # -- The walkthrough order from the methodology slide -------------------
    (4521, "Ananya Krishnan", "ananya.k@example.com", "9886112233",
     "Mixer Grinder", "Prestige Iris 750W", "2026-09-10", 2499.00,
     "Delivered", "Y", "21 Indiranagar, Bengaluru 560038", "TRK45210"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def _write_header(sheet, columns):
    sheet.append(columns)
    for index in range(1, len(columns) + 1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"


def _autosize(sheet, columns, cap=42):
    for index, name in enumerate(columns, start=1):
        widest = len(str(name))
        for cell in sheet[get_column_letter(index)]:
            if cell.value is not None:
                widest = max(widest, len(str(cell.value)))
        sheet.column_dimensions[get_column_letter(index)].width = min(widest + 2, cap)


def build(path=None):
    path = Path(path or config.CRM_WORKBOOK)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    orders = workbook.active
    orders.title = config.SHEET_ORDERS
    _write_header(orders, ORDER_COLUMNS)
    for row in ORDERS:
        orders.append(list(row))
    for row in orders.iter_rows(min_row=2, min_col=8, max_col=8):
        for cell in row:
            cell.number_format = '#,##0.00'
    _autosize(orders, ORDER_COLUMNS)

    # Both log sheets ship empty with headers only. The robot appends to them
    # at runtime, so a fresh workbook means a clean demo.
    audit = workbook.create_sheet(config.SHEET_AUDIT)
    _write_header(audit, AUDIT_COLUMNS)
    _autosize(audit, AUDIT_COLUMNS)

    escalations = workbook.create_sheet(config.SHEET_ESCALATION)
    _write_header(escalations, ESCALATION_COLUMNS)
    _autosize(escalations, ESCALATION_COLUMNS)

    workbook.save(path)
    return path


def main():
    path = build()
    print(f"Built mock CRM: {path}")
    print(f"  {config.SHEET_ORDERS:<16} {len(ORDERS)} orders across "
          f"{len(config.CATEGORIES)} appliance categories")
    print(f"  {config.SHEET_AUDIT:<16} empty (headers only)")
    print(f"  {config.SHEET_ESCALATION:<16} empty (headers only)")

    by_category = {}
    for row in ORDERS:
        by_category.setdefault(row[4], []).append(row)
    print()
    for category, rows in by_category.items():
        eligible = sum(1 for r in rows if r[9] == "Y")
        print(f"  {category:<18} {len(rows)} orders, {eligible} refund-eligible")


if __name__ == "__main__":
    main()
